"""모드 B 주변시설 PPTX (CLAUDE.md §5, /facilities/pptx).

FacilityResult → A3 1슬라이드: 위성 현황도(compose_map PNG 임베드 — 핀·반경원·범례)
+ 반경밴드별 개수표 + 시설목록 표. 표는 네이티브(편집가능)·새 숫자 0.
지도 라벨은 위성 PNG 에 구워지므로 한글 폰트가 필요(map_compose._font, Dockerfile fonts-nanum).
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

from app.schemas.facility import FacilityResult

_A3_W, _A3_H = Cm(42.0), Cm(29.7)
_NAVY = RGBColor(0x2F, 0x54, 0x96)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_F = "맑은 고딕"

# 시설목록 표에 담을 최대 행수 (거리순 상위) — 초과분은 캡션에 정직 표시
_LIST_MAX = 22


def _set_font(cell, size, bold=False, color=None, align=None):
    p = cell.text_frame.paragraphs[0]
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, _F
        if color is not None:
            run.font.color.rgb = color


def build_facilities_pptx(result: FacilityResult, radii: list[int], map_png: bytes | None) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = _A3_W, _A3_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── 제목 ──
    tf = slide.shapes.add_textbox(Cm(1.2), Cm(0.5), Cm(39), Cm(1.3)).text_frame
    tf.text = "주변 시설 현황"
    r0 = tf.paragraphs[0].runs[0]
    r0.font.size, r0.font.bold, r0.font.name, r0.font.color.rgb = Pt(22), True, _F, _NAVY

    tf2 = slide.shapes.add_textbox(Cm(1.2), Cm(1.6), Cm(39), Cm(0.9)).text_frame
    tf2.text = result.center.address
    if tf2.paragraphs[0].runs:
        tf2.paragraphs[0].runs[0].font.size = Pt(12)
        tf2.paragraphs[0].runs[0].font.name = _F

    # ── 왼쪽: 위성 현황도 (compose_map PNG 임베드 — 핀·반경원·범례 포함) ──
    ML, MT, MW = Cm(1.2), Cm(2.7), Cm(25.0)
    if map_png:
        try:
            slide.shapes.add_picture(io.BytesIO(map_png), ML, MT, MW, MW)
        except Exception:
            slide.shapes.add_textbox(ML, MT, MW, Cm(1)).text_frame.text = "(위성 지도 임베드 실패)"
    else:
        slide.shapes.add_textbox(ML, MT, MW, Cm(1)).text_frame.text = "(위성 지도 생성 실패)"

    # ── 오른쪽 상단: 반경밴드별 개수표 ──
    bands = sorted(result.counts.keys(), key=lambda x: int(x))
    kinds = [k for k in dict.fromkeys(f.kind for f in result.results)]
    if not kinds:  # 결과 0건이어도 counts 키로 종류 노출
        kinds = sorted({k for b in result.counts.values() for k in b})

    RL, RT, RW = Cm(27.0), Cm(2.7), Cm(13.8)
    n_rows = len(kinds) + 1
    ct = slide.shapes.add_table(n_rows, len(bands) + 1, RL, RT, RW,
                                Cm(0.9 * n_rows)).table
    ct.columns[0].width = Cm(4.8)
    for j in range(len(bands)):
        ct.columns[j + 1].width = Cm((13.8 - 4.8) / max(1, len(bands)))
    headers = ["종류"] + [f"{b}m" for b in bands]
    for j, h in enumerate(headers):
        cell = ct.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = _NAVY
        _set_font(cell, 11, bold=True, color=_WHITE, align=PP_ALIGN.CENTER)
    for i, k in enumerate(kinds, start=1):
        c0 = ct.cell(i, 0)
        c0.text = k
        _set_font(c0, 11)
        for j, b in enumerate(bands, start=1):
            cj = ct.cell(i, j)
            cj.text = str(result.counts.get(b, {}).get(k, 0))
            _set_font(cj, 11, bold=True, align=PP_ALIGN.CENTER)

    # ── 오른쪽 하단: 시설목록 (거리순 상위) ──
    items = sorted(result.results, key=lambda f: f.dist_m)[:_LIST_MAX]
    LT = RT + Cm(0.9 * n_rows) + Cm(0.6)
    if items:
        lt = slide.shapes.add_table(len(items) + 1, 3, RL, LT, RW,
                                    Cm(0.62 * (len(items) + 1))).table
        lt.columns[0].width, lt.columns[1].width, lt.columns[2].width = Cm(2.4), Cm(3.4), Cm(8.0)
        for j, h in enumerate(["거리", "종류", "시설명"]):
            cell = lt.cell(0, j)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = _NAVY
            _set_font(cell, 10, bold=True, color=_WHITE, align=PP_ALIGN.CENTER)
        for i, f in enumerate(items, start=1):
            d = lt.cell(i, 0); d.text = f"{f.dist_m}m"; _set_font(d, 9, align=PP_ALIGN.CENTER)
            kc = lt.cell(i, 1); kc.text = f.kind; _set_font(kc, 9, align=PP_ALIGN.CENTER)
            nc = lt.cell(i, 2); nc.text = f.name; _set_font(nc, 9)

    # ── 캡션 (출처·기준일·상한 정직 표시) ──
    extra = ""
    if len(result.results) > _LIST_MAX:
        extra = f" · 목록은 거리순 상위 {_LIST_MAX}건만 표기(전체 {len(result.results)}건, 개수표는 전량)"
    cap = slide.shapes.add_textbox(Cm(1.2), Cm(28.3), Cm(39), Cm(1.2)).text_frame
    cap.word_wrap = True
    cap.text = (f"배경: VWorld 항공영상 · 시설: {result.source} · 기준일 {result.base_date} · "
                f"반경(누적) {'·'.join(f'{b}m' for b in bands)}{extra}")
    if cap.paragraphs[0].runs:
        cap.paragraphs[0].runs[0].font.size, cap.paragraphs[0].runs[0].font.name = Pt(9), _F

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
