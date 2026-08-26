"""app/deck/render_slides.py — **밖에서 준 내용**을 우리 스타일로 그린다.

첫 사용처는 컨셉 스튜디오(`:8200`)의 컨셉 장표다. 그 앱이 「무엇을 적을 것인가」를 알고
우리가 「어떻게 그릴 것인가」를 안다(`deck_render/1.0`).

이 파일이 못박는 계약 셋:

- **우리는 「컨셉」이 무엇인지 모른다.** 아는 낱말은 `cover`·`kpi`·`cards`·`table`·`text`
  다섯뿐이고 전부 이미 우리 어휘다. 도메인 분기를 넣기 시작하면 deck-builder 가
  접힌 이유를 되풀이한다.
- **자른 것은 말한다.** 자리가 없어 뺀 줄을 조용히 없애면 발표자가 회의장에서 모른 채 넘어간다.
- **모르는 낱말도 장은 낸다.** 빈 장을 내면 왜 비었는지 아무도 모른다.

네트워크 없음 — 받은 dict 만으로 그린다.
"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Cm

import app.deck.render_slides as rs


def read(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def texts(slide) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]


def joined(slide) -> str:
    return " ".join(texts(slide))


DECK = {
    "title": "미아 레이어파크", "subtitle": "MIA LAYER PARK",
    "slides": [
        {"kind": "cover", "title": "미아 레이어파크", "subtitle": "MIA LAYER PARK",
         "lead": "25m를 세 켜로 접다", "body": ["mia130 · v8"],
         "sources": ["EV-0019"], "caption": "AI 제안 · 사람이 확정"},
        {"kind": "kpi", "title": "숫자", "kpis": [
            {"label": "총 세대수", "value": "1,730세대", "sub": "고시"},
            {"label": "공원 총량", "value": "12,180.2㎡", "sub": "파생"}],
         "sources": ["EV-0009", "EV-0064"]},
        {"kind": "cards", "title": "켜 / LAYER", "lead": "한 낱말의 세 갈래",
         "cards": [{"head": "쌓기", "body": "데크 3단", "tag": "겹쳐 올리다"},
                   {"head": "잇기", "body": "그린웨이 250m", "tag": "같은 높이를 잇다"}]},
        {"kind": "table", "title": "파트별", "headers": ["슬롯", "이름"],
         "rows": [["배치", "세 켜의 땅"], ["입면", "켜의 결"]], "ratios": [1.0, 2.0]},
    ],
}


# ── 그려진다 ────────────────────────────────────────────────────────────────


def test_it_draws_one_slide_per_spec():
    prs = read(rs.build(DECK))
    assert len(prs.slides) == len(DECK["slides"])


def test_the_page_is_a3():
    """`/deck/full` 과 같은 판이라야 한 회의에서 같이 쓴다."""
    prs = read(rs.build(DECK))
    assert (prs.slide_width, prs.slide_height) == (Cm(42.0), Cm(29.7))


def test_a_table_is_native_and_editable():
    """그림으로 박으면 회의 중에 못 고친다 — 그게 이 렌더러의 존재 이유다."""
    prs = read(rs.build(DECK))
    tables = [sh for sl in prs.slides for sh in sl.shapes if sh.has_table]
    assert len(tables) == 1
    tbl = tables[0].table
    assert len(tbl.rows) == 3 and len(tbl.columns) == 2      # 헤더 + 2줄
    assert tbl.cell(1, 1).text == "세 켜의 땅"


def test_the_cover_shows_name_and_slogan():
    prs = read(rs.build(DECK))
    body = joined(prs.slides[0])
    assert "미아 레이어파크" in body and "25m를 세 켜로 접다" in body


def test_every_card_makes_it_onto_the_slide():
    prs = read(rs.build(DECK))
    body = joined(prs.slides[2])
    assert "쌓기" in body and "잇기" in body and "그린웨이 250m" in body


def test_nothing_falls_off_the_page():
    """A3 밖으로 나간 도형은 인쇄하면 잘린다."""
    prs = read(rs.build(DECK))
    for n, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if None in (sh.left, sh.top, sh.width, sh.height):
                continue
            assert sh.left >= 0 and sh.top >= 0, f"{n}장 '{sh.name}' 이 판 왼쪽·위로 나갔다"
            assert sh.left + sh.width <= prs.slide_width + Cm(0.1), f"{n}장 '{sh.name}' 이 오른쪽으로"
            assert sh.top + sh.height <= prs.slide_height + Cm(0.1), f"{n}장 '{sh.name}' 이 아래로"


# ── 근거가 따라간다 ─────────────────────────────────────────────────────────


def test_the_evidence_numbers_are_printed():
    """PPTX 에는 `data-ev` 를 못 단다. 장표만 들고 심의에 가도 답할 수 있어야 한다."""
    prs = read(rs.build(DECK))
    assert "EV-0019" in joined(prs.slides[0])
    assert "EV-0009" in joined(prs.slides[1])


def test_a_long_evidence_list_is_summarised_not_dumped():
    """근거가 서른 개면 캡션이 한 줄을 넘는다 — 앞의 것만 찍고 나머지는 개수로."""
    deck = {"title": "x", "slides": [
        {"kind": "text", "title": "많은 근거", "body": ["본문"],
         "sources": [f"EV-{n:04d}" for n in range(1, 31)]}]}
    body = joined(read(rs.build(deck)).slides[0])
    assert "외 16건" in body


# ── 자른 것은 말한다 ────────────────────────────────────────────────────────


def test_extra_rows_are_cut_and_the_caption_says_so():
    deck = {"title": "x", "slides": [
        {"kind": "table", "title": "긴 표", "headers": ["가", "나"],
         "rows": [[str(n), "값"] for n in range(rs.MAX_ROWS + 5)]}]}
    prs = read(rs.build(deck))
    tbl = next(sh for sh in prs.slides[0].shapes if sh.has_table).table
    assert len(tbl.rows) == rs.MAX_ROWS + 1
    assert "5줄" in joined(prs.slides[0])


def test_extra_kpi_cards_are_cut_and_said():
    deck = {"title": "x", "slides": [
        {"kind": "kpi", "title": "많은 수치",
         "kpis": [{"label": f"항목{n}", "value": str(n)} for n in range(rs.MAX_KPI + 2)]}]}
    assert "2칸" in joined(read(rs.build(deck)).slides[0])


# ── 이상한 입력에도 장은 나온다 ─────────────────────────────────────────────


def test_an_unknown_kind_still_makes_a_slide():
    """빈 장을 내면 왜 비었는지 아무도 모른다 — 글로라도 남긴다."""
    deck = {"title": "x", "slides": [
        {"kind": "미래에생길것", "title": "모르는 낱말", "body": ["그래도 이 글은 남는다"]}]}
    prs = read(rs.build(deck))
    assert len(prs.slides) == 1
    assert "그래도 이 글은 남는다" in joined(prs.slides[0])


def test_an_empty_table_says_so_instead_of_crashing():
    deck = {"title": "x", "slides": [{"kind": "table", "title": "빈 표"}]}
    assert "표시할 표가 없다" in joined(read(rs.build(deck)).slides[0])


def test_ratios_that_do_not_match_the_columns_are_ignored():
    """열 수와 안 맞는 비율을 쓰면 표가 어긋난다 — 무시하고 기본 폭으로."""
    deck = {"title": "x", "slides": [
        {"kind": "table", "title": "어긋난 비율", "headers": ["가", "나", "다"],
         "rows": [["1", "2", "3"]], "ratios": [1.0, 2.0]}]}
    tbl = next(sh for sh in read(rs.build(deck)).slides[0].shapes if sh.has_table).table
    assert len(tbl.columns) == 3


def test_no_slides_is_refused_with_a_reason():
    """추정 대신 명확히 멈춘다(절대 원칙 3)."""
    with pytest.raises(ValueError):
        rs.build({"title": "x", "slides": []})


def test_it_does_not_need_to_know_what_a_concept_is():
    """우리가 아는 낱말은 다섯뿐이다. 늘면 이 앱이 남의 도메인을 알기 시작한 것이다."""
    assert set(rs._DRAWERS) == {"cover", "kpi", "cards", "table", "text"}
