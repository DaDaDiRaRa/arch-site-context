"""모드 B 주변시설 PPTX + 라우터 테스트 (CLAUDE.md §5, /facilities/pptx)."""

import io

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Cm

from app.main import app
import app.routers.facilities as frouter
from app.schemas.facility import Center, Facility, FacilityResult
from app.services.facilities_pptx import build_facilities_pptx

client = TestClient(app)


def _result():
    return FacilityResult(
        center=Center(lat=37.526, lon=126.9265, address="서울 영등포구 여의대로 24"),
        results=[
            Facility(kind="어린이집", name="여의도어린이집", lat=37.52, lon=126.92, dist_m=420, radius_band="500"),
            Facility(kind="경로당", name="여의경로당", lat=37.53, lon=126.93, dist_m=980, radius_band="1000"),
        ],
        counts={"500": {"어린이집": 3, "경로당": 5}, "1000": {"어린이집": 7, "경로당": 12}},
        source="kakao+vworld", base_date="2026-08-07")


def test_build_facilities_pptx_no_map():
    # 지도 실패해도 개수표·시설목록만으로 PPT 생성 (graceful)
    data = build_facilities_pptx(_result(), [500, 1000], None)
    assert data[:2] == b"PK"
    prs = Presentation(io.BytesIO(data))
    assert abs(prs.slide_width - Cm(42.0)) < 1000
    slides = list(prs.slides)
    assert len(slides) == 1
    s = slides[0]
    tables = [sh.table for sh in s.shapes if sh.has_table]
    assert len(tables) == 2  # 개수표 + 시설목록
    joined = " ".join(c.text for t in tables for r in t.rows for c in r.cells)
    assert "어린이집" in joined and "경로당" in joined
    assert "여의도어린이집" in joined and "여의경로당" in joined
    assert "500m" in joined and "1000m" in joined
    # 출처·기준일 캡션
    txt = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    assert "kakao" in txt and "2026-08-07" in txt


def test_facilities_pptx_endpoint(monkeypatch):
    monkeypatch.setattr(frouter, "build_facility_result", lambda *a, **k: _result())
    monkeypatch.setattr(frouter, "compose_map", lambda *a, **k: None)

    r = client.post("/facilities/pptx", json={
        "address": "서울 영등포구 여의대로 24", "kinds": ["어린이집", "경로당"], "radii": [500, 1000]})
    assert r.status_code == 200
    j = r.json()
    assert j["url"].endswith(".pptx")
    assert j["source"] == "kakao+vworld"
    assert j["counts"]["500"]["어린이집"] == 3
