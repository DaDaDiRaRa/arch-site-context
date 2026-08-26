"""app/deck/map_slides.py — 지도 슬라이드 4종(PPTX 네이티브 트랙)의 graceful 계약.

`map_svg.py`(SVG 트랙)는 처음부터 fetch 실패를 방어했지만 PPTX 원본은 반환값을 그대로
언패킹해, 위성타일이 한 번 실패하면 `/deck/full` 전체가 500 이었다. 여기서 그 계약을 못박는다:
지도 하나가 죽어도 덱은 나머지로 완결된다(절대 원칙 3).

네트워크는 전부 monkeypatch — 실호출 0.
"""

from __future__ import annotations

import io

from pptx import Presentation
from PIL import Image

import app.deck.clients as clients
import app.deck.map_slides as ms


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (90, 110, 130)).save(buf, format="PNG")
    return buf.getvalue()


FAKE_META = {"zoom": 17, "cx": 1000.0, "cy": 1000.0, "radius_px": 500.0, "size_px": 1500}


# ── 지목 파싱 (VWorld jibun = 지번+지목) ─────────────────────────────────────

def test_parse_jimok_strips_any_lot_number() -> None:
    """지번은 주소마다 다르다 — 특정 지번을 하드코딩하면 그 주소에서만 맞는다."""
    assert ms.parse_jimok("385대") == "대"
    assert ms.parse_jimok("24대") == "대"        # 예전엔 '24대' 가 지목 칸에 그대로 나왔다
    assert ms.parse_jimok("24-1대") == "대"
    assert ms.parse_jimok("산12임") == "임"
    assert ms.parse_jimok("3850대") == "대"


def test_parse_jimok_blank_when_unparseable() -> None:
    """못 믿을 값은 추정하지 않고 비운다 → 호출부가 '확인필요' 로 (절대 원칙 3)."""
    assert ms.parse_jimok("385") == ""            # 지목 없음
    assert ms.parse_jimok("") == ""
    assert ms.parse_jimok(None) == ""
    assert ms.parse_jimok("385대 12") == ""       # 숫자가 남으면 지목으로 못 믿음
    assert ms.parse_jimok("서울시 영등포구") == ""  # 지목 부호는 한 글자 — 긴 값은 안 앉힌다


# ── 지도 슬라이드 graceful ───────────────────────────────────────────────────

def test_slide_wide_skips_when_basemap_fails(monkeypatch) -> None:
    """fetch_basemap 은 실패 시 None — 언패킹하면 TypeError 로 덱 전체가 죽는다."""
    monkeypatch.setattr(clients, "fetch_basemap", lambda *a, **k: None)
    monkeypatch.setattr(clients, "fetch_facilities", lambda *a, **k: None)
    prs = Presentation()
    assert ms.slide_wide(prs, "서울 영등포구 여의대로 24", 37.52, 126.93, {}, {}) is False


def test_slide_wide_survives_facilities_none(monkeypatch) -> None:
    """시설 조회만 실패해도(지도는 살아있음) 슬라이드는 만들어진다."""
    monkeypatch.setattr(clients, "fetch_basemap", lambda *a, **k: (FAKE_META, _png()))
    monkeypatch.setattr(clients, "fetch_facilities", lambda *a, **k: None)
    prs = Presentation()
    prs.slide_width, prs.slide_height = ms.k.A3_W, ms.k.A3_H
    ms.slide_wide(prs, "서울 영등포구 여의대로 24", 37.52, 126.93, {}, {})
    assert len(prs.slides) == 1


def test_match_names_tolerates_missing_facilities() -> None:
    blds = [{"clat": 37.52, "clon": 126.93, "name": None}]
    ms._match_names(blds, None)  # 예전엔 AttributeError
    assert blds[0]["name"] is None


def test_full_deck_survives_total_map_failure(monkeypatch) -> None:
    """외부 소스가 전부 죽어도 덱은 나온다 — 라우터가 500 이 아니라 PPTX 를 준다."""
    class _Site:
        lat, lon, pnu = 37.52, 126.93, "1156010100103850000"

    monkeypatch.setattr("app.services.site_seed.build_site", lambda addr: _Site())
    for name in ("fetch_basemap", "fetch_model", "fetch_law", "fetch_site",
                 "fetch_board", "fetch_surroundings", "fetch_facilities"):
        monkeypatch.setattr(clients, name, lambda *a, **k: None)

    data = ms.build_full_deck("서울 영등포구 여의대로 24")
    assert data.startswith(b"PK")  # 유효한 pptx(zip)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 1  # 최소한 커버는 남는다
